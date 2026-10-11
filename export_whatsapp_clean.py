#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Export WhatsApp Outbox to a Clean, Deduplicated Excel File (.xlsx)
Send-off ya Lilian Marcus Nyahende
"""

import os
import json
import re
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

def normalize_phone(phone_str):
    if not phone_str:
        return ""
    digits = re.sub(r'[^0-9]', '', str(phone_str))
    if digits.startswith('0') and len(digits) == 10:
        digits = '255' + digits[1:]
    elif len(digits) == 9:
        digits = '255' + digits
    return digits

def load_database():
    live_path = os.path.join(BASE_DIR, 'data', 'live_db.json')
    db_path = os.path.join(BASE_DIR, 'data', 'db.json')
    
    file_to_use = live_path if os.path.exists(live_path) else db_path
    if not os.path.exists(file_to_use):
        raise FileNotFoundError(f"Database file not found at {live_path} or {db_path}")
    
    print(f"[*] Inasoma database kutoka: {file_to_use}")
    with open(file_to_use, 'r', encoding='utf-8') as f:
        return json.load(f)

def clean_and_deduplicate(db):
    logs = db.get('whatsappLogs', [])
    guests = db.get('guests', [])

    print(f"[*] Jumla ya logs kwenye database: {len(logs)}")
    print(f"[*] Jumla ya wageni kwenye database: {len(guests)}")

    # Index guests by normalized phone and by clean lower name
    guest_by_phone = {}
    guest_by_name = {}
    for g in guests:
        p = normalize_phone(g.get('phone', ''))
        if p:
            guest_by_phone[p] = g
        n = (g.get('name') or '').strip().lower()
        if n:
            guest_by_name[n] = g

    # Deduplicate logs by unique phone (or clean name if phone missing)
    recipients = {}

    for log in logs:
        p = normalize_phone(log.get('recipientPhone') or log.get('phone') or '')
        name = (log.get('recipientName') or log.get('name') or '').strip()
        key = p if p else name.lower()

        if not key:
            continue

        status = str(log.get('status', '')).lower()
        is_delivered = (status == 'delivered')
        timestamp = log.get('timestamp') or log.get('time') or ''
        provider = log.get('provider') or 'Meta Cloud'

        # Match extra details from guest registry
        matched_g = guest_by_phone.get(p) or (guest_by_name.get(name.lower()) if name else None)
        final_name = name or (matched_g.get('name') if matched_g else 'Mualikwa')
        seats_val = matched_g.get('seats', 1) if matched_g else 1
        is_double = (str(seats_val).lower() in ['2', 'double', 'double vip'] or matched_g.get('isDouble') is True) if matched_g else False
        table_num = (matched_g.get('table') or '') if matched_g else ''

        if key not in recipients:
            recipients[key] = {
                'name': final_name,
                'phone': p,
                'seats': 'Double (2)' if is_double else 'Single (1)',
                'table': str(table_num) if table_num else '-',
                'status': 'delivered' if is_delivered else 'failed',
                'success_count': 1 if is_delivered else 0,
                'failed_count': 0 if is_delivered else 1,
                'attempts': 1,
                'provider': provider,
                'last_timestamp': timestamp
            }
        else:
            item = recipients[key]
            item['attempts'] += 1
            if is_delivered:
                item['status'] = 'delivered'  # Prioritize delivered!
                item['success_count'] += 1
                item['provider'] = provider
            else:
                item['failed_count'] += 1

            if timestamp and (not item['last_timestamp'] or timestamp > item['last_timestamp']):
                item['last_timestamp'] = timestamp

    all_list = list(recipients.values())
    all_list.sort(key=lambda x: x['name'].lower())

    delivered_list = [r for r in all_list if r['status'] == 'delivered']
    failed_list = [r for r in all_list if r['status'] == 'failed']

    print(f"[OK] Wageni wa kipekee waliotumiwa (Delivered): {len(delivered_list)}")
    print(f"[FAIL] Wageni waliokwama kabisa (Failed): {len(failed_list)}")
    print(f"[TOTAL] Jumla ya wageni wa kipekee: {len(all_list)}")

    return delivered_list, failed_list, all_list

def create_excel_report(delivered_list, failed_list, all_list, output_filename):
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styles
    navy_header_fill = PatternFill(start_color="0B2240", end_color="0B2240", fill_type="solid")
    gold_header_fill = PatternFill(start_color="CBA135", end_color="CBA135", fill_type="solid")
    green_badge_fill = PatternFill(start_color="E6F4EA", end_color="E6F4EA", fill_type="solid")
    red_badge_fill = PatternFill(start_color="FCE8E6", end_color="FCE8E6", fill_type="solid")
    
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=10, bold=True)
    font_regular = Font(name="Calibri", size=10)
    font_green = Font(name="Calibri", size=10, bold=True, color="137333")
    font_red = Font(name="Calibri", size=10, bold=True, color="C5221F")

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    
    thin_border = Border(
        left=Side(style='thin', color='DDDDDD'),
        right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'),
        bottom=Side(style='thin', color='DDDDDD')
    )

    # ========================================================
    # SHEET 1: WALIOTUMIWA (IMEFIKA) - 100% CLEAN
    # ========================================================
    ws1 = wb.create_sheet(title="Waliotumiwa (Imefika)")
    ws1.views.sheetView[0].showGridLines = True

    # Title Banner
    ws1.merge_cells("A1:H1")
    title_cell = ws1["A1"]
    title_cell.value = "ORODHA YA WAGENI WALIOTUMIWA KADI YA WHATSAPP (SEND-OFF YA LILIAN MARCUS)"
    title_cell.font = Font(name="Calibri", size=13, bold=True, color="0B2240")
    title_cell.alignment = align_left
    ws1.row_dimensions[1].height = 25

    headers_1 = [
        "Na.", "Jina la Mgeni", "Namba ya WhatsApp", "Viti", 
        "Hali ya Utumaji", "Mtoa Huduma", "Muda Uliotumwa", "Maelezo ya Utumaji"
    ]

    for col_idx, header in enumerate(headers_1, 1):
        cell = ws1.cell(row=3, column=col_idx, value=header)
        cell.fill = navy_header_fill
        cell.font = font_header
        cell.alignment = align_center
        cell.border = thin_border
    ws1.row_dimensions[3].height = 24

    for row_idx, r in enumerate(delivered_list, start=4):
        phone_display = f"+{r['phone']}" if r['phone'] else "-"
        note = "Ilifanikiwa baada ya kurudiwa" if r['failed_count'] > 0 else "Ilifika mara ya kwanza"
        
        row_values = [
            row_idx - 3,
            r['name'],
            phone_display,
            r['seats'],
            "IMEFIKA",
            r['provider'],
            r['last_timestamp'][:16].replace('T', ' ') if r['last_timestamp'] else "-",
            note
        ]

        for col_idx, val in enumerate(row_values, 1):
            cell = ws1.cell(row=row_idx, column=col_idx, value=val)
            cell.font = font_regular
            cell.border = thin_border
            if col_idx in [1, 3, 4, 6, 7]:
                cell.alignment = align_center
            else:
                cell.alignment = align_left

            if col_idx == 5:  # Status
                cell.font = font_green
                cell.fill = green_badge_fill
                cell.alignment = align_center

        ws1.row_dimensions[row_idx].height = 20

    # Auto-fit columns
    for col in ws1.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws1.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # ========================================================
    # SHEET 2: WALIOKWAMA (FAILED)
    # ========================================================
    ws2 = wb.create_sheet(title="Waliokwama (Hawakupokea)")
    ws2.views.sheetView[0].showGridLines = True

    ws2.merge_cells("A1:G1")
    t2 = ws2["A1"]
    t2.value = "ORODHA YA WAGENI WAMBAO UJUMBE WA WHATSAPP ULIKWAMA (HAUKUFIKA)"
    t2.font = Font(name="Calibri", size=13, bold=True, color="C5221F")
    t2.alignment = align_left
    ws2.row_dimensions[1].height = 25

    headers_2 = [
        "Na.", "Jina la Mgeni", "Namba ya WhatsApp", "Viti", 
        "Hali ya Utumaji", "Mtoa Huduma", "Hatua Inayopendekezwa"
    ]

    for col_idx, header in enumerate(headers_2, 1):
        cell = ws2.cell(row=3, column=col_idx, value=header)
        cell.fill = PatternFill(start_color="8A1B14", end_color="8A1B14", fill_type="solid")
        cell.font = font_header
        cell.alignment = align_center
        cell.border = thin_border
    ws2.row_dimensions[3].height = 24

    for row_idx, r in enumerate(failed_list, start=4):
        phone_display = f"+{r['phone']}" if r['phone'] else "-"
        row_values = [
            row_idx - 3,
            r['name'],
            phone_display,
            r['seats'],
            "IMEKWAMA",
            r['provider'],
            "Piga simu au tuma SMS ya kawaida ya NextSMS"
        ]

        for col_idx, val in enumerate(row_values, 1):
            cell = ws2.cell(row=row_idx, column=col_idx, value=val)
            cell.font = font_regular
            cell.border = thin_border
            if col_idx in [1, 3, 4, 6]:
                cell.alignment = align_center
            else:
                cell.alignment = align_left

            if col_idx == 5:  # Status
                cell.font = font_red
                cell.fill = red_badge_fill
                cell.alignment = align_center

        ws2.row_dimensions[row_idx].height = 20

    for col in ws2.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws2.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # ========================================================
    # SHEET 3: MUHTASARI WA TAKWIMU
    # ========================================================
    ws3 = wb.create_sheet(title="Muhtasari wa Takwimu")
    ws3.views.sheetView[0].showGridLines = True

    ws3.merge_cells("A1:C1")
    t3 = ws3["A1"]
    t3.value = "RIPOTI NA MUHTASARI WA UTUMAJI WA KADI ZA WHATSAPP"
    t3.font = Font(name="Calibri", size=13, bold=True, color="0B2240")
    ws3.row_dimensions[1].height = 25

    headers_3 = ["Kipengele / Maelezo", "Takwimu", "Asilimia (%)"]
    for col_idx, header in enumerate(headers_3, 1):
        cell = ws3.cell(row=3, column=col_idx, value=header)
        cell.fill = gold_header_fill
        cell.font = Font(name="Calibri", size=11, bold=True, color="0B2240")
        cell.alignment = align_center
        cell.border = thin_border

    total_unique = len(all_list)
    total_delivered = len(delivered_list)
    total_failed = len(failed_list)
    pct_delivered = f"{(total_delivered / total_unique * 100):.1f}%" if total_unique else "0%"
    pct_failed = f"{(total_failed / total_unique * 100):.1f}%" if total_unique else "0%"

    stats = [
        ("Jumla ya Wageni wa Kipekee Waliojaribiwa", total_unique, "100.0%"),
        ("Waliopokea Kadi kwa Mafanikio (Delivered)", total_delivered, pct_delivered),
        ("Waliokwama / Hawajapokea (Failed)", total_failed, pct_failed),
        ("Tarehe ya Kuandaa Ripoti", datetime.now().strftime("%d/%m/%Y %H:%M"), "-"),
    ]

    for row_idx, (lbl, val, pct) in enumerate(stats, start=4):
        ws3.cell(row=row_idx, column=1, value=lbl).font = font_bold
        ws3.cell(row=row_idx, column=2, value=val).font = font_regular
        ws3.cell(row=row_idx, column=3, value=pct).font = font_regular
        for c in range(1, 4):
            ws3.cell(row=row_idx, column=c).border = thin_border
            if c > 1:
                ws3.cell(row=row_idx, column=c).alignment = align_center

    ws3.column_dimensions["A"].width = 45
    ws3.column_dimensions["B"].width = 20
    ws3.column_dimensions["C"].width = 18

    wb.save(output_filename)
    print(f"[SUCCESS] Faili la Excel limehifadhiwa kwa mafanikio: {output_filename}")

if __name__ == "__main__":
    db = load_database()
    delivered, failed, all_recipients = clean_and_deduplicate(db)
    
    date_str = datetime.now().strftime("%Y_%m_%d")
    out_file = os.path.join(BASE_DIR, f"Wageni_Waliotumiwa_WhatsApp_Safi_{date_str}.xlsx")
    create_excel_report(delivered, failed, all_recipients, out_file)
